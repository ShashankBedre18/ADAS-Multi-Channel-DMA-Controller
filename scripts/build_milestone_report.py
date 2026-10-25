"""
build_milestone_report.py - writes docs/ADAS_DMA_Milestones_Report.docx
(milestone-by-milestone report with diagrams; figures come from docs/figures).
Needs: pip install python-docx
"""
import os, re
from docx import Document
from docx.shared import Pt, Inches, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = os.path.normpath(os.path.join(os.path.dirname(__file__), ".."))
FIG = os.path.join(ROOT, "docs", "figures")
OUT = os.path.join(ROOT, "docs", "ADAS_DMA_Milestones_Report.docx")

NAVY = RGBColor(0x1E, 0x3A, 0x5F)
BLUE = RGBColor(0x2F, 0x5D, 0x9A)
GREY = RGBColor(0x5B, 0x64, 0x72)

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(2.2)
sec.top_margin, sec.bottom_margin = Cm(2.0), Cm(2.0)
TEXT_W = 21.0 - 4.4  # cm

st = doc.styles["Normal"]
st.font.name = "Calibri"
st.font.size = Pt(11)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
st.paragraph_format.space_after = Pt(6)
st.paragraph_format.line_spacing = 1.12
for lvl, size, col in ((1, 18, NAVY), (2, 14, BLUE), (3, 12, NAVY)):
    h = doc.styles[f"Heading {lvl}"]
    h.font.name = "Calibri"
    h.element.rPr.rFonts.set(qn("w:asciiTheme"), "") if False else None
    h.font.size = Pt(size)
    h.font.bold = True
    h.font.color.rgb = col
    h.paragraph_format.space_before = Pt(16 if lvl == 1 else 12)
    h.paragraph_format.space_after = Pt(6)
    rf = h.element.rPr.find(qn("w:rFonts"))
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if rf is not None and rf.get(qn(a)) is not None:
            del rf.attrib[qn(a)]
    if rf is not None:
        rf.set(qn("w:ascii"), "Calibri"); rf.set(qn("w:hAnsi"), "Calibri")

fig_no = [0]
tab_no = [0]


# ---------------------------------------------------------------- helpers
def shade(cell, hex_):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), hex_)
    tcPr.append(sh)


def cell_margins(table, top=60, bottom=60, left=100, right=100):
    tblPr = table._tbl.tblPr
    m = OxmlElement("w:tblCellMar")
    for k, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        e = OxmlElement(f"w:{k}"); e.set(qn("w:w"), str(v)); e.set(qn("w:type"), "dxa"); m.append(e)
    look = tblPr.find(qn("w:tblLook"))
    if look is not None:
        look.addprevious(m)
    else:
        tblPr.append(m)


def runs(p, text, size=None, color=None, bold_all=False):
    """**bold**, `code` inline markup."""
    for part in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
        if not part:
            continue
        if part.startswith("**"):
            r = p.add_run(part[2:-2]); r.bold = True
        elif part.startswith("`"):
            r = p.add_run(part[1:-1]); r.font.name = "Consolas"; r.font.size = Pt((size or 11) - 1)
            r.font.color.rgb = RGBColor(0x9B, 0x2C, 0x2C)
        else:
            r = p.add_run(part)
        if bold_all:
            r.bold = True
        if size and not part.startswith("`"):
            r.font.size = Pt(size)
        if color:
            r.font.color.rgb = color
    return p


def P(text, align=None, size=None, italic=False, after=None):
    p = doc.add_paragraph()
    runs(p, text, size)
    if italic:
        for r in p.runs: r.italic = True
    if align == "j":
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    elif align == "c":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    return p


def B(items, style="List Bullet"):
    for it in items:
        p = doc.add_paragraph(style=style)
        runs(p, it)
        p.paragraph_format.space_after = Pt(3)


def N(items):
    for i, it in enumerate(items, 1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.9)
        p.paragraph_format.first_line_indent = Cm(-0.6)
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(f"{i}.  "); r.bold = True; r.font.color.rgb = BLUE
        runs(p, it)


TOC_ENTRIES = []


def H(text, lvl):
    if lvl <= 2 and text != "Contents":
        TOC_ENTRIES.append((lvl, text))
    return doc.add_heading(text, lvl)


def T(header, rows, widths, hdr_fill="1E3A5F", zebra=True, size=9.5, caption=None):
    if caption:
        tab_no[0] += 1
        c = doc.add_paragraph()
        r = c.add_run(f"Table {tab_no[0]}. {caption}"); r.bold = True; r.font.size = Pt(9.5); r.font.color.rgb = NAVY
        c.paragraph_format.space_after = Pt(2); c.paragraph_format.keep_with_next = True
    t = doc.add_table(rows=1 + len(rows), cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    cell_margins(t)
    for i, w in enumerate(widths):
        for row in t.rows:
            row.cells[i].width = Cm(w)
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]; c.text = ""
        r = c.paragraphs[0].add_run(h); r.bold = True; r.font.size = Pt(size); r.font.color.rgb = RGBColor(255, 255, 255)
        shade(c, hdr_fill)
    for ri, row in enumerate(rows):
        for ci, v in enumerate(row):
            c = t.rows[ri + 1].cells[ci]; c.text = ""
            runs(c.paragraphs[0], str(v), size)
            c.paragraphs[0].paragraph_format.space_after = Pt(0)
            if zebra and ri % 2 == 1:
                shade(c, "F3F6FA")
    # repeat header row
    trPr = t.rows[0]._tr.get_or_add_trPr()
    e = OxmlElement("w:tblHeader"); e.set(qn("w:val"), "true"); trPr.append(e)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def CODE(text, title=None):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell_margins(t, 80, 80, 140, 140)
    c = t.rows[0].cells[0]; c.width = Cm(TEXT_W)
    shade(c, "F4F5F7")
    c.text = ""
    first = True
    if title:
        r = c.paragraphs[0].add_run(title); r.bold = True; r.font.size = Pt(9); r.font.color.rgb = GREY
        first = False
    for ln in text.strip("\n").split("\n"):
        p = c.paragraphs[0] if first else c.add_paragraph()
        first = False
        p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.0
        r = p.add_run(ln if ln else " "); r.font.name = "Consolas"; r.font.size = Pt(8.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def BOX(title, lines, fill="E8F0FB", edge="2F5D9A"):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell_margins(t, 100, 100, 160, 160)
    c = t.rows[0].cells[0]; c.width = Cm(TEXT_W)
    shade(c, fill)
    tcPr = c._tc.get_or_add_tcPr()
    b = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        e = OxmlElement(f"w:{side}"); e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "8"); e.set(qn("w:color"), edge)
        if side == "left": e.set(qn("w:sz"), "24")
        b.append(e)
    sh = tcPr.find(qn("w:shd"))
    if sh is not None:
        sh.addprevious(b)
    else:
        tcPr.append(b)
    c.text = ""
    r = c.paragraphs[0].add_run(title); r.bold = True; r.font.color.rgb = NAVY
    for ln in lines:
        p = c.add_paragraph(); p.paragraph_format.space_after = Pt(2)
        runs(p, ln, 10.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def FIGURE(name, caption, width=TEXT_W):
    fig_no[0] += 1
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(os.path.join(FIG, name), width=Cm(width))
    c = doc.add_paragraph(); c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = c.add_run(f"Figure {fig_no[0]}. "); r.bold = True; r.font.size = Pt(9.5); r.font.color.rgb = NAVY
    r = c.add_run(caption); r.italic = True; r.font.size = Pt(9.5); r.font.color.rgb = GREY
    c.paragraph_format.space_after = Pt(10)


def ASK(milestone_item, built, evidence):
    """standard header block for each milestone point"""
    t = doc.add_table(rows=3, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = "Table Grid"
    cell_margins(t, 60, 60, 120, 120)
    for i, (k, v, f) in enumerate((("What the milestone asks", milestone_item, "FDF0E2"),
                                    ("What we built", built, "E7F6EC"),
                                    ("Evidence (files)", evidence, "F1F3F5"))):
        a, b = t.rows[i].cells
        a.width, b.width = Cm(4.2), Cm(TEXT_W - 4.2)
        a.text = ""; r = a.paragraphs[0].add_run(k); r.bold = True; r.font.size = Pt(9.5); r.font.color.rgb = NAVY
        shade(a, f); shade(b, f)
        b.text = ""; runs(b.paragraphs[0], v, 9.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def PAGEBREAK():
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def field(par, instr):
    r = par.add_run()
    f1 = OxmlElement("w:fldChar"); f1.set(qn("w:fldCharType"), "begin")
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = instr
    f2 = OxmlElement("w:fldChar"); f2.set(qn("w:fldCharType"), "separate")
    r._r.append(f1); r._r.append(it); r._r.append(f2)
    return par


# ---------------------------------------------------------------- cover
for _ in range(3):
    doc.add_paragraph()
p = P("L&T EduTech  |  VLSI Principles, RTL Design and Verification", "c", 12); p.runs[0].font.color.rgb = GREY
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Project 1"); r.font.size = Pt(16); r.bold = True; r.font.color.rgb = BLUE
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Multi-Channel DMA Controller Subsystem"); r.font.size = Pt(28); r.bold = True; r.font.color.rgb = NAVY
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("for an Automotive ADAS Platform"); r.font.size = Pt(18); r.font.color.rgb = NAVY
doc.add_paragraph()
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Milestone-wise Project Report"); r.font.size = Pt(15); r.bold = True
P("Milestone 1: Requirements and Architecture Design", "c", 12, after=2)
P("Milestone 2: RTL Implementation", "c", 12, after=2)
P("Milestone 3: Verification and Validation", "c", 12, after=14)
doc.add_paragraph().paragraph_format.space_after = Pt(24)
T(["Item", "Details"], [
    ["Student name", ""], ["Registration / LMS ID", ""],
    ["Language and methodology", "SystemVerilog RTL, SystemVerilog Assertions, UVM 1.2"],
    ["Tools", "AMD Vivado 2025.2 (xsim simulator, xcrg coverage, synthesis)"],
    ["Target", "4-channel DMA, APB control + AXI4 data, 100 MHz (Artix-7 xc7a100t)"],
    ["Final results", "13/13 tests pass, 100 % functional coverage, 100 % RTL statement coverage"],
], [5.0, 11.0], zebra=False, size=10)
PAGEBREAK()

# ---------------------------------------------------------------- TOC
H("Contents", 1)
TOC_PAR = field(doc.add_paragraph(), 'TOC \\o "1-2" \\h \\z \\u')
P("(In Word: right-click the list and choose Update Field to add page numbers.)", size=9, italic=True)
TOC_ANCHOR = doc.paragraphs[-1]
PAGEBREAK()

# ================================================================ 1. INTRODUCTION
H("1. Introduction", 1)
H("1.1 The problem statement given in the brief", 2)
P("A leading automotive SoC manufacturer is building a next-generation **Advanced Driver Assistance System (ADAS)** "
  "platform. The platform must continuously move high-bandwidth sensor data from **cameras, radar and LiDAR** "
  "into memory while keeping CPU overhead low. In the existing design the processor itself moves the data "
  "(processor-centric data movement). This creates performance bottlenecks and extra latency during real-time "
  "operation. The company therefore needs a **configurable DMA subsystem** that can:", "j")
B(["manage **multiple concurrent data streams**,",
   "**prioritise critical sensor transfers** (for example radar, which feeds emergency braking),",
   "support **burst transactions** on the memory bus,",
   "generate **status interrupts**, and",
   "provide **reliable error handling**."])
FIGURE("m1_cpu_vs_dma.png", "Why a DMA is needed: on the left the CPU copies every sensor word; on the right the DMA moves the data in bursts and the CPU only programs registers and handles interrupts.")

H("1.2 What we were asked to build, and what we delivered", 2)
P("Every requirement in the brief is mapped below to the feature that satisfies it and the file where it is implemented. "
  "The rest of this report follows the three milestones exactly, point by point.", "j")
T(["Requirement in the brief", "How the design meets it", "Where"], [
    ["Multiple concurrent data streams", "4 independent channels; each can be memory-to-memory or sensor-stream-to-memory; a shared burst engine serves them one burst at a time", "rtl/dma_channel.sv, rtl/dma_engine.sv"],
    ["Prioritise critical sensor transfers", "2-bit priority per channel (3 = most critical) + round-robin between equal priorities", "rtl/dma_arbiter.sv"],
    ["Burst transactions", "AXI4 INCR bursts of 1 to 16 beats, automatically split so that no burst crosses a 4 KB page", "rtl/dma_channel.sv (sizing), rtl/dma_engine.sv"],
    ["Status interrupts", "Per-channel DONE / ERR interrupt bits (write-1-to-clear) + mask + global irq line", "rtl/dma_channel.sv, rtl/dma_top.sv"],
    ["Reliable error handling", "AXI SLVERR / DECERR capture, illegal configuration check, restart after an error; an error stops only the affected channel", "rtl/dma_engine.sv, rtl/dma_channel.sv"],
], [4.2, 8.4, 4.0], caption="Brief requirements mapped to the design")

H("1.3 Verification requirements of the brief", 2)
T(["Brief says", "What we did", "Section"], [
    ["Develop a SystemVerilog transaction-based testbench", "UVM testbench: transaction items (APB, sensor, AXI), drivers, monitors, interfaces with clocking blocks", "4.1, 4.2"],
    ["Develop assertions to verify ...", "25 SVA properties (52 instances) for APB, AXI handshake, AXI legality incl. 4 KB rule, arbitration order, channel and IRQ rules", "4.4"],
    ["Create covergroups to measure ...", "7 covergroups: configuration, AXI bursts and responses, arbitration, interrupts, errors, sensor back-pressure, APB access", "4.5"],
    ["Generate functional coverage reports", "Merged HTML reports from Vivado xcrg: 100 % functional, 100 % statement coverage on every RTL module", "4.5"],
], [5.0, 9.0, 2.6], caption="Verification requirements of the brief")

H("1.4 The design at a glance", 2)
T(["Parameter", "Value"], [
    ["Number of channels", "4 (CH0 camera, CH1 radar, CH2 LiDAR, CH3 diagnostics in the ADAS scenario)"],
    ["Control interface", "APB slave, 12-bit address, 32-bit data, zero wait states"],
    ["Data interface", "AXI4 master, 32-bit address and data, INCR bursts, 1-16 beats, single outstanding burst"],
    ["Sensor inputs", "4 valid/ready word streams, each with a 32-word FIFO"],
    ["Transfer length", "1 to 16 M words per channel (24-bit LEN)"],
    ["Transfer modes", "memory-to-memory or sensor-to-memory"],
    ["Arbitration", "fixed priority (0-3) + round-robin among equals, one decision per burst"],
    ["Interrupts", "irq_ch[3:0] + global irq; causes DONE and ERR"],
    ["Errors detected", "SLVERR, DECERR, illegal configuration"],
    ["Clock / reset", "single clock, asynchronous active-low reset, 100 MHz timing met"],
    ["RTL size", "6 SystemVerilog files, about 690 lines, fully synthesizable"],
], [4.5, 12.1], caption="Key specifications")

H("1.5 How this report is organised", 2)
T(["Milestone", "Items in the brief", "Report sections"], [
    ["Milestone 1: Requirements and Architecture Design", "1. Study DMA architecture used in SoCs\n2. Define channel configuration registers\n3. Design arbitration scheme\n4. Create datapath and controller architecture\n5. Develop RTL block diagram", "2.1 - 2.5"],
    ["Milestone 2: RTL Implementation", "1. Implement DMA channel controller\n2. Design arbitration logic\n3. Implement transfer state machine\n4. Integrate interrupt and status registers\n5. Perform RTL simulation", "3.1 - 3.5"],
    ["Milestone 3: Verification and Validation", "1. Develop SystemVerilog testbench\n2. Create UVM agents and sequences\n3. Implement scoreboarding and checking\n4. Add assertions for protocol compliance\n5. Generate functional and code coverage reports", "4.1 - 4.5"],
], [4.6, 9.0, 3.0], caption="Milestones and where each item is answered")
P("Each milestone item starts with a small box: **what the milestone asks**, **what we built**, and the **files** that prove it.", italic=True)
PAGEBREAK()

# ================================================================ 2. MILESTONE 1
H("2. Milestone 1: Requirements and Architecture Design", 1)
P("Milestone 1 turns the problem into a precise specification: how DMA engines work in SoCs, what the software "
  "sees (registers), how competing requests are resolved (arbitration), how data and control flow (datapath and "
  "controller), and the resulting block diagram. Nothing is coded yet; every decision made here is implemented in Milestone 2.", "j")

# ---- 2.1
H("2.1 Study DMA architecture used in SoCs", 2)
ASK("Study how DMA controllers are built and used inside SoCs.",
    "A study of DMA concepts, buses and design options, and a requirements list (R1-R10) derived from the ADAS problem.",
    "docs/01_problem_and_requirements.md, docs/02_architecture.md")
H("2.1.1 What a DMA controller is", 3)
P("A **Direct Memory Access (DMA)** controller is a small hardware block that copies data from a source to a destination "
  "without the CPU touching each word. The CPU only writes a few registers (source, destination, length, mode) and starts "
  "the channel. The DMA then becomes a **bus master**, reads and writes memory on its own using efficient bursts, and "
  "raises an **interrupt** when the job is finished or when something went wrong.", "j")
H("2.1.2 Where a DMA sits in an SoC", 3)
P("A typical SoC has two kinds of buses. A slow, simple **control bus (APB)** is used by the CPU to program peripherals. "
  "A fast **data bus (AXI4)** connects bus masters (CPU, DMA, accelerators) to memory. A DMA is therefore a **slave on APB** "
  "(it is programmed) and a **master on AXI** (it moves data).", "j")
FIGURE("m1_soc_context.png", "System context: the DMA subsystem is programmed over APB, streams data from the sensor interfaces and writes it to DDR through the AXI4 interconnect.")
H("2.1.3 Building blocks found in SoC DMA controllers", 3)
T(["Building block", "Purpose", "Used in this design"], [
    ["Channels", "Independent jobs, each with its own source/destination/length", "Yes, 4 channels"],
    ["Register-based or descriptor-based setup", "Registers = simple; descriptors in memory = long scatter-gather chains", "Register-based (simple, deterministic)"],
    ["Arbiter", "Chooses which channel uses the bus next", "Yes, priority + round-robin"],
    ["Burst engine", "Turns a transfer into legal bus bursts", "Yes, AXI4 INCR, 1-16 beats"],
    ["Data buffer / FIFO", "Holds data between read and write, absorbs rate differences", "16-word burst buffer + 32-word FIFO per sensor"],
    ["Interrupt / status logic", "Tells software about completion and errors", "Yes, DONE / ERR per channel"],
    ["Error handling", "Bus errors, illegal settings", "SLVERR, DECERR, ERR_CFG"],
], [4.0, 7.2, 5.4], caption="Typical DMA building blocks and our choices")
H("2.1.4 Bus protocol essentials", 3)
P("**AXI4** has five independent channels: read address (AR), read data (R), write address (AW), write data (W) and write "
  "response (B). Every channel uses a **VALID / READY handshake**: a transfer happens on a clock edge where both are high, and "
  "the sender must keep VALID and the payload stable until READY arrives. Key rules used in this design:", "j")
B(["**INCR bursts**: one address, then `AxLEN+1` beats at consecutive addresses (we use 1 to 16 beats of 4 bytes).",
   "**4 KB rule**: a burst must never cross a 4 KB address boundary.",
   "**WLAST** marks the last write beat; **RLAST** marks the last read beat.",
   "**Responses**: OKAY (00), SLVERR (10, slave error) and DECERR (11, no slave at that address)."])
P("**APB** is a two-phase protocol: a SETUP cycle (PSEL=1, PENABLE=0) followed by an ACCESS cycle (PENABLE=1) that completes when "
  "PREADY=1. PSLVERR flags an illegal access.", "j")
H("2.1.5 One DMA transfer, step by step", 3)
N(["Software writes SRC, DST, LEN and INT_EN of a channel, then writes CTRL with EN=1 and START=1.",
   "The channel checks the configuration. If it is illegal it raises ERR (code 3) and stays idle.",
   "The channel becomes active and computes the size of its next burst.",
   "When it has a burst ready, it raises a request to the arbiter.",
   "The arbiter grants the bus to the most important requester.",
   "The engine reads one burst (from memory over AXI, or from the sensor FIFO) into the burst buffer.",
   "The engine writes the burst to the destination over AXI and waits for the write response.",
   "The channel advances its addresses and remaining count; steps 3-7 repeat.",
   "After the last word the channel sets DONE and the interrupt line rises. Software reads and clears INT_STAT."])
H("2.1.6 ADAS sensors and the requirements derived from them", 3)
T(["Sensor", "Behaviour", "What matters", "Our setting"], [
    ["Front camera", "very high bandwidth, continuous frames", "throughput, no lost pixels", "CH0, priority 1, 16-beat bursts"],
    ["Radar", "small bursty packets", "latency (emergency braking)", "CH1, priority 3, 4-beat bursts"],
    ["LiDAR", "medium-bandwidth point clouds", "throughput + bounded latency", "CH2, priority 2"],
    ["Diagnostics", "bulk memory-to-memory copy", "must not disturb the sensors", "CH3, priority 0"],
], [2.7, 4.2, 4.2, 5.5], caption="ADAS sensor characteristics")
T(["ID", "Requirement"], [
    ["R1", "4 independent channels, each memory-to-memory or sensor-stream-to-memory."],
    ["R2", "AXI4 memory interface: 32-bit, INCR bursts up to 16 beats, never crossing a 4 KB boundary."],
    ["R3", "Channels configured through an APB register interface."],
    ["R4", "A higher-priority channel is served before lower ones at every burst boundary."],
    ["R5", "Equal-priority channels share the bus fairly (round-robin)."],
    ["R6", "Per-channel interrupt on completion and on error; maskable, write-1-to-clear."],
    ["R7", "SLVERR/DECERR stops only the affected channel and records the cause; others continue."],
    ["R8", "Illegal configurations are rejected at START (ERR_CFG) with no bus traffic."],
    ["R9", "After an error a channel can be restarted by software."],
    ["R10", "Data is never lost or reordered; a sensor that cannot be accepted is back-pressured."],
], [1.3, 15.3], caption="Functional requirements (all verified in Milestone 3)")
H("2.1.7 Architecture decisions taken from the study", 3)
T(["Decision", "Options considered", "Chosen", "Reason"], [
    ["Setup method", "registers / descriptors in memory", "registers", "simple, deterministic, enough for 4 streams"],
    ["Bus engine", "one per channel / shared", "one shared engine", "one AXI master port, smaller area, arbitration point is clear"],
    ["Data movement", "cut-through / store-and-forward", "store-and-forward (16 words)", "a write burst never stalls waiting for data, AXI rules trivially met"],
    ["Arbitration point", "per word / per burst / per transfer", "per burst", "critical channel waits at most one burst"],
    ["Tie-break", "fixed order / round-robin", "round-robin", "equal channels share the bus fairly"],
], [3.0, 4.6, 3.6, 5.4], caption="Design choices")

# ---- 2.2
H("2.2 Define channel configuration registers", 2)
ASK("Define the registers software uses to configure each channel.",
    "A complete APB register map: 4 global registers and 9 registers per channel, with bit fields, access types, reset values and hardware-enforced rules.",
    "docs/03_register_map.md, rtl/dma_pkg.sv (offsets), rtl/dma_channel.sv, rtl/dma_top.sv")
P("The register space is 4 KB (12-bit APB address). Global registers live at the bottom; each channel has its own 64-byte "
  "window starting at `0x100 + n x 0x40`. Any other address returns **PSLVERR**, so software bugs are caught immediately.", "j")
FIGURE("m1_address_map.png", "APB address map: global registers, four channel windows and the registers inside one window.")
FIGURE("m1_ctrl_bits.png", "Bit fields of CTRL, STATUS and INT_STAT.")
T(["Offset", "Name", "Access", "Reset", "Fields / meaning"], [
    ["0x000", "GCTRL", "RW", "0", "[0] EN global enable"],
    ["0x004", "GSTAT", "RO", "0", "[3:0] channel active, [8] engine busy"],
    ["0x008", "IRQ_SUM", "RO", "0", "[3:0] per-channel interrupt lines"],
    ["0x00C", "VERSION", "RO", "0x00010000", "design version"],
], [1.6, 2.4, 1.5, 2.1, 9.0], caption="Global registers")
T(["Offset", "Name", "Access", "Fields / meaning"], [
    ["+0x00", "CTRL", "RW", "[0] EN, [1] START (pulse), [2] SRC_STREAM (1 = sensor), [5:4] PRIO, [11:8] BLEN = beats-1"],
    ["+0x04", "SRC", "RW", "source byte address (word aligned; ignored for sensor streams)"],
    ["+0x08", "DST", "RW", "destination byte address (word aligned)"],
    ["+0x0C", "LEN", "RW", "transfer length in words [23:0]"],
    ["+0x10", "INT_EN", "RW", "[0] DONE, [1] ERR interrupt enables"],
    ["+0x14", "INT_STAT", "W1C", "[0] DONE, [1] ERR flags; write 1 to clear"],
    ["+0x18", "STATUS", "RO", "[0] ACTIVE, [1] IN_FLIGHT, [4:2] ERR_CODE, [13:8] FIFO_LEVEL"],
    ["+0x1C", "REMAIN", "RO", "words left in the current transfer"],
    ["+0x20", "XFERRED", "RO", "words moved since START"],
], [1.6, 2.4, 1.5, 11.1], caption="Channel registers (offset from the channel base)")
P("**ERR_CODE values:** 0 none, 1 SLVERR, 2 DECERR, 3 illegal configuration.")
H("Rules the hardware enforces", 3)
B(["START while the channel is already active is ignored.",
   "START with EN=0, LEN=0, an unaligned DST, or (memory source) an unaligned SRC sets ERR with ERR_CODE=3 and causes **no bus traffic**.",
   "After an error the channel is idle; software restarts it with a new START.",
   "A new START clears ERR_CODE and XFERRED and flushes the channel FIFO.",
   "INT_STAT: if a new event and a software clear happen in the same cycle, the event wins (no lost interrupt)."])
CODE("""write(GCTRL,  1);                       // global enable
// radar: sensor stream, priority 3, 4-beat bursts, 256 words
write(CH1_DST,    0x80100000);
write(CH1_LEN,    256);
write(CH1_INT_EN, 0b11);                // DONE | ERR
write(CH1_CTRL,   (3<<8)|(3<<4)|(1<<2)|(1<<1)|1);   // BLEN=3 PRIO=3 STREAM START EN
// ... wait for irq ...
st = read(CH1_INT_STAT);  write(CH1_INT_STAT, st);  // acknowledge (W1C)""", "Programming example (what a driver does)")

# ---- 2.3
H("2.3 Design arbitration scheme for multiple DMA requests", 2)
ASK("Design how the DMA decides which of several requesting channels uses the bus.",
    "A per-burst arbiter: static 2-bit priority, round-robin between equal priorities.",
    "docs/02_architecture.md section 2.5, rtl/dma_arbiter.sv")
H("2.3.1 What the arbiter must guarantee", 3)
B(["**Criticality**: radar (priority 3) must be served before camera and diagnostics.",
   "**Bounded latency**: a critical channel must never wait for a whole long transfer of another channel.",
   "**Fairness**: channels with the same priority share the bus."])
T(["Scheme", "Advantage", "Problem for ADAS"], [
    ["Fixed priority with fixed tie order", "simple, radar always first", "equal channels are not treated fairly"],
    ["Round-robin only", "perfectly fair", "radar waits behind camera; criticality ignored"],
    ["Weighted round-robin", "bandwidth shares", "more logic, latency of radar not strictly bounded"],
    ["**Priority + round-robin (chosen)**", "criticality first, fairness among equals", "very small: one 2-bit pointer"],
], [4.6, 5.6, 6.4], caption="Arbitration schemes compared")
H("2.3.2 The chosen scheme", 3)
N(["Each channel presents **req** (active, enabled and - for sensors - enough data in its FIFO) and a 2-bit **PRIO**.",
   "The requester with the highest PRIO wins.",
   "Ties are broken by a **round-robin pointer** that starts just after the last winner.",
   "Arbitration happens **once per burst**, so a critical channel waits at most one burst (at most 16 beats) of another channel."])
FIGURE("m1_arb_flow.png", "Arbitration decision flow, executed every time the engine is ready for a new burst.")
FIGURE("m1_arb_example.png", "Worked example. Top: different priorities - the channels are served strictly in priority order. Bottom: equal priorities - the grant rotates (round-robin).")
H("2.3.3 Worst-case latency of the critical channel", 3)
P("When radar requests, the engine may be in the middle of a burst of another channel. The longest burst is 16 beats. "
  "A 16-beat memory-to-memory burst takes roughly 16 read beats + 16 write beats + address, response and the two pipeline "
  "states, i.e. about **40 to 45 clock cycles (about 0.45 us at 100 MHz)** with a zero-wait memory. After that burst radar is "
  "granted immediately, because it has the highest priority. This bound is independent of the "
  "transfer length of the other channels - exactly the property that the processor-centric design lacked.", "j")

# ---- 2.4
H("2.4 Create datapath and controller architecture", 2)
ASK("Define how data moves (datapath) and which blocks make the decisions (controller).",
    "A store-and-forward datapath through one 16-word burst buffer, and a three-level controller: channel controllers, arbiter, engine FSM.",
    "docs/02_architecture.md sections 2.2-2.4, rtl/dma_engine.sv, rtl/dma_channel.sv")
H("2.4.1 Datapath", 3)
P("Every burst passes through a single **16 x 32-bit burst buffer** inside the engine. For a memory source the buffer is "
  "filled by an AXI read burst; for a sensor source it is filled from that channel's FIFO. The buffer is then emptied by an "
  "AXI write burst to the destination.", "j")
FIGURE("m1_datapath.png", "Datapath: sources, multiplexer, burst buffer and AXI write path (top); address and burst-size path (bottom).")
P("**Why store-and-forward?** The write burst starts only when all its data is already in the buffer, so WVALID never has to "
  "drop in the middle of a burst and WLAST is trivially correct. The bus is free between bursts so the arbiter can switch "
  "channels. The cost is one burst of latency, which is why bursts are limited to 16 beats and the burst length is "
  "programmable per channel (radar uses 4 beats for lower latency, the camera 16 for higher efficiency).", "j")
H("2.4.2 Burst sizing and the 4 KB rule", 3)
P("Before each burst the channel computes the number of beats:")
CODE("beats = min( BLEN + 1,                         // programmed burst size\n"
     "             remaining words,                 // do not overrun the transfer\n"
     "             words left in destination 4 KB page,\n"
     "             words left in source 4 KB page )  // memory source only")
P("**Example.** DST = 0x0000_0FEC (5 words below a 4 KB boundary), LEN = 40, BLEN = 15. Bursts: "
  "**5** (reach the boundary) + **16** + **16** + **3** = 40 words. No burst crosses 0x1000, so the AXI rule is met "
  "without any software effort.", "j")
H("2.4.3 Controller architecture", 3)
FIGURE("m1_controller.png", "Controller split: channel controllers decide IF they need the bus, the arbiter decides WHO gets it, the engine decides HOW the burst runs.")
T(["Signal", "From -> To", "Meaning"], [
    ["req[c], prio[c]", "channel -> arbiter", "channel c has a burst ready; its static priority"],
    ["next_beats[c], src_addr[c], dst_addr[c]", "channel -> engine", "parameters of the next burst"],
    ["gnt_valid, gnt_idx", "arbiter -> engine", "a winner exists and which channel it is"],
    ["grant_fire", "engine -> arbiter", "engine accepts the grant (updates rr_ptr)"],
    ["fifo_pop[c], fifo_rdata[c]", "engine <-> channel", "read sensor words out of the channel FIFO"],
    ["inflight[c]", "engine -> channel", "engine is currently serving channel c"],
    ["upd[c], upd_beats, upd_err", "engine -> channel", "burst finished: how many words moved and the error code"],
    ["irq_ch[c]", "channel -> top", "interrupt of channel c; ORed into irq"],
], [5.0, 3.6, 8.0], caption="Interfaces between the controller blocks")

# ---- 2.5
H("2.5 Develop RTL block diagram", 2)
ASK("Draw the RTL block diagram of the subsystem.",
    "The top-level block diagram, the module hierarchy and the full port list of dma_top.",
    "docs/block_diagram.svg, rtl/dma_top.sv")
FIGURE("block_diagram.png", "RTL block diagram of the ADAS multi-channel DMA subsystem.")
FIGURE("m1_hierarchy.png", "Module hierarchy: one top, four channel instances (each with a FIFO), one arbiter, one engine, shared package.")
T(["Port group", "Signals", "Direction"], [
    ["Clock / reset", "clk, rst_n (asynchronous, active low)", "in"],
    ["APB slave", "psel, penable, pwrite, paddr[11:0], pwdata[31:0] / prdata[31:0], pready, pslverr", "in / out"],
    ["Sensor streams (x4)", "s_valid[3:0], s_data[3:0][31:0] / s_ready[3:0]", "in / out"],
    ["AXI write address", "awaddr, awlen, awsize, awburst, awvalid / awready", "out / in"],
    ["AXI write data", "wdata, wstrb, wlast, wvalid / wready", "out / in"],
    ["AXI write response", "bresp, bvalid / bready", "in / out"],
    ["AXI read address", "araddr, arlen, arsize, arburst, arvalid / arready", "out / in"],
    ["AXI read data", "rdata, rresp, rlast, rvalid / rready", "in / out"],
    ["Interrupts", "irq_ch[3:0], irq", "out"],
], [3.6, 10.0, 3.0], caption="dma_top ports")
T(["Parameter (dma_pkg)", "Value", "Meaning"], [
    ["N_CH", "4", "number of channels"], ["DATA_W / ADDR_W", "32 / 32", "bus widths"],
    ["MAX_BEATS", "16", "longest burst"], ["FIFO_DEPTH", "32", "words per sensor FIFO"], ["LEN_W", "24", "width of LEN"],
], [4.6, 2.6, 9.4], caption="Design parameters")
BOX("Milestone 1 - outcome", [
    "All five items are complete: SoC DMA study and requirements (2.1), register map (2.2), arbitration scheme (2.3), "
    "datapath and controller architecture (2.4), and RTL block diagram (2.5).",
    "Every later RTL module and every test traces back to a requirement R1-R10 defined here."], "E7F6EC", "2F855A")
PAGEBREAK()

# ================================================================ 3. MILESTONE 2
H("3. Milestone 2: RTL Implementation", 1)
P("Milestone 2 implements the architecture in synthesizable SystemVerilog. The design is split into the six modules "
  "shown in Figure 10. All of them use one clock and an asynchronous active-low reset, and all constants come from `dma_pkg`.", "j")
T(["Module", "Lines", "Role"], [
    ["dma_pkg.sv", "54", "parameters, error codes, register offsets, AXI constants"],
    ["dma_sfifo.sv", "44", "32-word first-word-fall-through FIFO for one sensor (maps to LUT-RAM)"],
    ["dma_channel.sv", "199", "channel controller: registers, context, burst sizing, request, interrupts"],
    ["dma_arbiter.sv", "47", "priority + round-robin arbiter"],
    ["dma_engine.sv", "173", "transfer state machine, burst buffer, AXI4 master"],
    ["dma_top.sv", "169", "APB decoder, global registers, instances, irq OR"],
], [3.6, 1.6, 11.4], caption="RTL files")

# ---- 3.1
H("3.1 Implement DMA channel controller", 2)
ASK("Implement the logic that holds one channel's configuration and drives its transfer.",
    "dma_channel: configuration registers, START handling with configuration check, run-time context, sensor FIFO, burst sizing, registered request, interrupt and status logic. Instantiated 4 times.",
    "rtl/dma_channel.sv, rtl/dma_sfifo.sv")
FIGURE("m2_channel_internal.png", "Inside one channel controller.")
H("3.1.1 Responsibilities", 3)
B(["**Configuration registers** written over APB: EN, PRIO, BLEN, SRC_STREAM, SRC, DST, LEN, INT_EN.",
   "**START check**: a START is accepted only if EN=1, LEN != 0 and the addresses are word aligned; otherwise ERR_CFG.",
   "**Run-time context**: `cur_src`, `cur_dst`, `remaining`, `xferred`, `err_code`; updated after every burst.",
   "**Sensor FIFO**: accepts words while `s_ready = active & stream & !full`; a full FIFO holds the sensor off (back-pressure).",
   "**Burst sizing** with the min() formula of section 2.4.2.",
   "**Request**: raised when the channel has work and, for sensors, enough words in the FIFO for the next burst.",
   "**Completion**: on the last burst it raises DONE and goes idle; on a bus error it raises ERR and goes idle."])
FIGURE("m2_channel_fsm.png", "Channel life-cycle.")
CODE("""wire cfg_bad   = !new_en || (len == '0) ||
                 (dst_base[1:0] != 2'b00) ||
                 (!new_strm && src_base[1:0] != 2'b00);
wire start_ok  = start_w && !active && !cfg_bad;   // accepted START
wire start_bad = start_w && !active &&  cfg_bad;   // rejected -> ERR_CFG, no bus traffic

// after every burst the engine calls back with upd / upd_beats / upd_err
if (upd) begin
  cur_src   <= cur_src + (upd_beats << 2);
  cur_dst   <= cur_dst + (upd_beats << 2);
  remaining <= remaining - upd_beats;
  if (upd_err != ERR_NONE)          begin active <= 0; err_code <= upd_err; end  // error: stop
  else if (remaining == upd_beats)  active <= 0;                                // last burst
end""", "dma_channel.sv - START check and burst completion (simplified excerpt)")
P("**Timing detail.** The request and next-burst size are **registered** (`req_q`, `nb_q`). This removed a long "
  "combinational path (context -> min() -> arbiter -> engine) found during synthesis; the engine waits one SETTLE cycle after "
  "each burst so that the registered request is up to date.", "j")

# ---- 3.2
H("3.2 Design arbitration logic", 2)
ASK("Implement the arbitration scheme of section 2.3 in RTL.",
    "dma_arbiter: a rotating scan that finds the highest-priority requester, and a 2-bit round-robin pointer for ties.",
    "rtl/dma_arbiter.sv")
FIGURE("m2_arbiter_hw.png", "Arbiter hardware structure.")
CODE("""// winner: highest priority, ties broken by round-robin order
for (int i = 0; i < N_CH; i++) begin
  idx = (rr_ptr + i) % N_CH;
  if (req[idx] && (!gnt_valid || prio[idx] > best)) begin
    gnt_valid = 1; gnt_idx = idx; best = prio[idx];
  end
end

// next tie starts after the last winner
always_ff @(posedge clk or negedge rst_n)
  if (!rst_n)          rr_ptr <= '0;
  else if (grant_fire) rr_ptr <= (gnt_idx == N_CH-1) ? '0 : gnt_idx + 1'b1;""", "dma_arbiter.sv - core of the arbiter (excerpt)")
P("Because the scan starts at `rr_ptr` and only replaces the current best with a **strictly higher** priority, the first "
  "channel after the previous winner wins any tie. This gives round-robin among equals with no extra logic.", "j")
FIGURE("wave_priority.png", "Simulation of dma_priority_test, part 1: four channels with priorities E4h (ch3=3, ch2=2, ch1=1, ch0=0) request at the same time; gnt_idx serves them strictly in the order 3, 2, 1, 0.")
FIGURE("wave_roundrobin.png", "Simulation of dma_priority_test, part 2: four channels with equal priority (AAh); at every grant_fire gnt_idx and rr_ptr rotate 1, 2, 3, 0, ... - round-robin. The test checks that every 4 consecutive grants cover all 4 channels.")

# ---- 3.3
H("3.3 Implement transfer state machine", 2)
ASK("Implement the FSM that performs the bus transfers.",
    "dma_engine: a 10-state FSM that runs one burst per pass (read or sensor pop, then write), the burst buffer, the AXI4 master signals and error capture.",
    "rtl/dma_engine.sv")
FIGURE("m2_engine_fsm.png", "Transfer state machine of dma_engine.")
T(["State", "What happens", "Leaves when"], [
    ["IDLE", "waits; grant_fire = gnt_valid; remembers the winning channel", "a grant is valid"],
    ["LATCH", "copies beats, source and destination address of the winner", "next cycle"],
    ["RD_AR", "ARVALID = 1 with ARADDR/ARLEN", "ARREADY"],
    ["RD_DATA", "RREADY = 1, stores each beat in the buffer, keeps the first non-OKAY response", "RVALID & RLAST"],
    ["STR_RD", "pops `beats` words from the channel FIFO into the buffer", "last word popped"],
    ["WR_AW", "AWVALID = 1 with AWADDR/AWLEN", "AWREADY"],
    ["WR_W", "WVALID = 1, WDATA = buffer[cnt], WLAST on the last beat", "WREADY & WLAST"],
    ["WR_B", "BREADY = 1, records BRESP", "BVALID"],
    ["UPD", "tells the channel how many words moved and the error code (a read error skips the write)", "next cycle"],
    ["SETTLE", "one idle cycle so the channel's registered request is up to date", "next cycle"],
], [2.4, 10.4, 3.8], caption="Engine states")
H("3.3.1 Error handling", 3)
B(["The engine keeps the first non-OKAY response of a burst in an `err` register.",
   "A SLVERR/DECERR during the read phase skips the write (no corrupted data reaches memory); during the write phase it is reported from BRESP.",
   "On any error the channel moves 0 words (`upd_beats = 0`), stops and records ERR_CODE; other channels continue (R7)."])
CODE("""case (state)
  S_IDLE:    if (gnt_valid)        nstate = S_LATCH;
  S_LATCH:                         nstate = src_stream[cur] ? S_STR_RD : S_RD_AR;
  S_RD_AR:   if (arready)          nstate = S_RD_DATA;
  S_RD_DATA: if (rvalid && rlast)  nstate = (read error) ? S_UPD : S_WR_AW;
  S_STR_RD:  if (cnt == beats-1)   nstate = S_WR_AW;
  S_WR_AW:   if (awready)          nstate = S_WR_W;
  S_WR_W:    if (wready && wlast)  nstate = S_WR_B;
  S_WR_B:    if (bvalid)           nstate = S_UPD;
  S_UPD:                           nstate = S_SETTLE;
  S_SETTLE:                        nstate = S_IDLE;
endcase""", "dma_engine.sv - next-state logic (excerpt)")
FIGURE("wave_axi_burst.png", "Simulation of one memory-to-memory burst: AR handshake, 16 read beats ending with RLAST, AW handshake, 16 write beats ending with WLAST, then the B response. Next burst starts at address +0x40.")
P("**Reading the waveform.** `engine_state` goes RD_AR -> RD_DATA (16 beats of `rdata`) -> WR_AW -> WR_W (16 beats of "
  "`wdata`, `wlast` on the last) -> WR_B (`bvalid`, `bresp = 0` OKAY) -> UPD -> SETTLE -> next burst. `arlen = awlen = F` means 16 beats, "
  "and the second burst starts at 0x10040 / 0x100040 (16 words x 4 bytes = 0x40 later).", "j")

# ---- 3.4
H("3.4 Integrate interrupt and status registers", 2)
ASK("Add interrupt and status registers and connect them to the controller.",
    "Per-channel INT_STAT (W1C) / INT_EN, STATUS, REMAIN, XFERRED; global GSTAT and IRQ_SUM; irq_ch[3:0] and irq outputs.",
    "rtl/dma_channel.sv, rtl/dma_top.sv")
FIGURE("m2_irq_logic.png", "Interrupt logic: two event sources set INT_STAT; software clears it by writing 1; INT_EN masks it; the channel lines are ORed into irq.")
T(["Bit", "Name", "Set when", "Typical software reaction"], [
    ["0", "DONE", "all LEN words moved", "process the buffer, start the next transfer"],
    ["1", "ERR", "SLVERR, DECERR or bad configuration; cause in STATUS.ERR_CODE", "log, fix, restart the channel"],
], [1.0, 1.8, 7.4, 6.4], caption="Interrupt causes")
CODE("""assign int_set = { (upd && upd_err != ERR_NONE) || start_bad,   // ERR
                   done_now };                                  // DONE
// set has priority over a simultaneous write-1-to-clear
int_stat <= (int_stat & ~(wr_int_stat ? reg_wdata[1:0] : 2'b00)) | int_set;
assign irq = |(int_stat & int_en);     // per channel; top: irq = |irq_ch""", "dma_channel.sv - interrupt status (excerpt)")
P("**Status visibility.** STATUS gives ACTIVE, IN_FLIGHT, ERR_CODE and the FIFO level; REMAIN and XFERRED give progress. "
  "Globally, GSTAT shows which channels are active and whether the engine is busy, and IRQ_SUM shows all channel "
  "interrupt lines in one read.", "j")
FIGURE("wave_sensor_stream.png", "Simulation of a sensor transfer on channel 1: words enter the FIFO (s_valid / s_ready), the engine drains them as 8-beat bursts to consecutive addresses, and irq_ch[1] and irq rise when the transfer is complete.")

# ---- 3.5
H("3.5 Perform RTL simulation", 2)
ASK("Simulate the RTL to show that it works.",
    "A scripted compile / elaborate / simulate flow in Vivado xsim, 8 directed tests + constrained-random tests with 5 seeds, waveform dumps, logs and coverage databases. FPGA synthesis as an extra check.",
    "sim/ (build, run_test, regress, coverage_report scripts), sim/logs/, synth/")
FIGURE("m2_sim_flow.png", "Simulation and regression flow.")
CODE("""sim\\regress.bat                              :: build + all 13 runs + coverage reports
sim\\run_test.bat dma_sensor_fusion_test       :: one test   (args: test [seed] [iterations])
xsim tb_snap -testplusarg UVM_TESTNAME=dma_smoke_test -testplusarg WAVES -R   :: with VCD""", "Commands")
FIGURE("regression.png", "Regression result of the final run: 13 of 13 runs pass, 0 UVM errors, 0 fatals.", 13.0)
H("3.5.1 The ADAS scenario in simulation", 3)
P("`dma_sensor_fusion_test` runs all four streams at the same time:")
T(["Channel", "Words", "Finished after (cycles)", "Observation"], [
    ["CH1 radar (P3)", "256", "1 120", "finishes first although it started together with the camera (checked by the test)"],
    ["CH2 LiDAR (P2)", "1 024", "4 277", "served before camera and diagnostics"],
    ["CH0 camera (P1)", "4 608", "15 978", "crosses a 4 KB page; bursts split automatically"],
    ["CH3 diagnostics (P0)", "512", "17 345", "lowest priority: uses the bus when the sensors are idle"],
], [3.8, 1.8, 3.6, 7.4], caption="Sensor-fusion scenario")
P("All 6 400 words were checked by the scoreboard; the completion order follows the priorities exactly.")
H("3.5.2 Synthesis check (beyond the brief)", 3)
T(["Metric", "Result (xc7a100t, Vivado 2025.2, out-of-context)"], [
    ["LUTs", "1 660 (2.6 %)"], ["Flip-flops", "1 065 (0.8 %)"], ["BRAM / DSP", "0 / 0"],
    ["Clock", "100 MHz, WNS +1.742 ns, all constraints met, 0 critical warnings (post-synthesis estimate)"],
], [4.0, 12.6], caption="Synthesis results")
P("During development the first synthesis attempt failed timing (WNS -8.1 ns) and used far too many flip-flops. Two RTL fixes solved it: memories "
  "without reset (so they map to LUT-RAM) and registering the channel request (the LATCH and SETTLE states).", "j")
BOX("Milestone 2 - outcome", [
    "All five items are complete: channel controller (3.1), arbitration logic (3.2), transfer state machine (3.3), interrupt and "
    "status registers (3.4), RTL simulation with passing regression (3.5).",
    "The RTL is synthesizable and meets 100 MHz on an Artix-7 FPGA."], "E7F6EC", "2F855A")
PAGEBREAK()

# ================================================================ 4. MILESTONE 3
H("4. Milestone 3: Verification and Validation", 1)
P("Milestone 3 proves that the RTL meets requirements R1-R10. The approach is a **self-checking, coverage-driven UVM "
  "testbench**: stimulus is generated as transactions, a reference model predicts what the DUT must do, assertions check "
  "protocol rules every clock cycle, and covergroups measure which scenarios have really been exercised.", "j")

# ---- 4.1
H("4.1 Develop SystemVerilog testbench", 2)
ASK("Develop a SystemVerilog transaction-based testbench.",
    "tb_top with clock/reset, four interface types with clocking blocks, the DUT, bound assertions, and transaction classes for every bus.",
    "tb/tb_top.sv, tb/if/dma_ifs.sv, tb/uvm/dma_tb_items.svh")
FIGURE("m3_tb_arch.png", "Testbench top level.")
T(["Interface", "Connects", "Clocking blocks"], [
    ["apb_if", "APB master side of the DUT", "drv_cb (drive), mon_cb (sample)"],
    ["axi_if", "AXI4 slave side (memory model)", "slv_cb (respond), mon_cb (sample)"],
    ["strm_if x4", "one per sensor stream", "drv_cb, mon_cb"],
    ["dbg_if", "hierarchical probe of arbiter req/grant", "mon_cb"],
], [3.0, 7.0, 6.6], caption="Testbench interfaces")
P("Clocking blocks drive outputs 1 ns after the clock edge and sample inputs just before it (`#1step`), so the testbench "
  "never races with the DUT.", "j")
T(["Transaction", "Fields", "Used by"], [
    ["apb_item", "write, addr, data, delay", "APB driver / monitor"],
    ["strm_item", "nwords, gap_pct, start_delay", "sensor driver (random gaps model a real sensor)"],
    ["strm_mon_item", "channel, word, accepted / stalled", "stream monitor -> scoreboard, coverage"],
    ["axi_item", "direction, address, length, data[], response", "AXI monitor -> scoreboard, coverage"],
    ["arb_item", "requests, priorities, winner", "arbitration monitor -> coverage"],
], [3.2, 6.4, 7.0], caption="Transaction classes")

# ---- 4.2
H("4.2 Create UVM agents and sequences", 2)
ASK("Create UVM agents (driver, monitor, sequencer) and sequences for the stimulus.",
    "3 agent types (APB, sensor stream x4, reactive AXI memory slave), an arbitration monitor, a virtual sequencer and 9 virtual sequences / tests.",
    "tb/uvm/dma_tb_agents.svh, dma_tb_seqs.svh, dma_tb_tests.svh, dma_tb_env.svh")
FIGURE("m3_uvm.png", "UVM environment.")
FIGURE("m3_seq_flow.png", "Life of one transaction.")
T(["Agent", "Components", "Role"], [
    ["apb_agent", "sequencer, driver, monitor", "programs registers like the CPU would; monitor reports every access"],
    ["strm_agent x4", "sequencer, driver, monitor", "plays camera / radar / LiDAR / diag data with random gaps; monitor reports accepted and stalled words"],
    ["axi_agent", "reactive slave driver, monitor", "memory model; random latency, back-pressure, SLVERR / DECERR injection"],
    ["arb_monitor", "monitor", "records every arbitration decision for coverage"],
], [3.0, 4.2, 9.4], caption="Agents")
T(["Test / sequence", "Scenario", "Requirements"], [
    ["dma_regs_test", "reset values, RW/RO, LEN mask, unmapped -> PSLVERR, W1C, interrupt lines", "R3, R6"],
    ["dma_smoke_test", "one memory transfer and one sensor transfer", "R1, R2"],
    ["dma_sensor_fusion_test", "the ADAS scenario: camera, radar, LiDAR and diag together; radar must finish first", "R1, R4, R10"],
    ["dma_burst_sweep_test", "every BLEN 0..15, random lengths, addresses just below a 4 KB page", "R2"],
    ["dma_priority_test", "strict priority order; long high-priority job before short low-priority job; equal priority rotation", "R4, R5"],
    ["dma_error_test", "SLVERR read, DECERR write, error on a sensor channel, recovery, 4x4 error matrix", "R6, R7, R9"],
    ["dma_cfg_err_test", "zero length, unaligned addresses, START without EN, START while busy", "R8"],
    ["dma_backpressure_test", "very slow memory, three fast sensors: FIFOs fill, sensors are held off, no data lost", "R10"],
    ["dma_random_test", "constrained-random mix of everything; 5 seeds x 40 iterations", "all"],
], [4.0, 9.8, 2.8], caption="Tests (each runs one virtual sequence)")

# ---- 4.3
H("4.3 Implement scoreboarding and checking", 2)
ASK("Implement a scoreboard that checks the DUT automatically.",
    "dma_scoreboard with an independent per-channel reference model built only from observed APB writes; seven classes of checks; a mutation experiment proving the checks work.",
    "tb/uvm/dma_tb_env.svh (ch_model, dma_scoreboard)")
FIGURE("m3_scoreboard.png", "Scoreboard structure and checks.")
P("The model does **not** look inside the RTL. It re-derives the expected burst sizes with its own copy of the min() "
  "formula and its own 4 KB arithmetic, and it remembers the source data (a memory snapshot at START, or the exact words "
  "each sensor sent) to compare with every beat written to memory. Any difference is reported as a UVM_ERROR and fails the test.", "j")
T(["Injected bug (one at a time)", "Caught by"], [
    ["data bit flipped on beat 5 of a burst", "dma_smoke_test: scoreboard DATA MISMATCH word 5"],
    ["arbiter ignores priority", "dma_priority_test: strict priority order check"],
    ["4 KB splitting removed", "dma_burst_sweep_test: scoreboard burst-length check"],
    ["DECERR reported as SLVERR", "dma_error_test: ERR_CODE check (got 1, expected 2)"],
], [7.0, 9.6], caption="Mutation check on the final RTL - each injected bug made its test FAIL")

# ---- 4.4
H("4.4 Add assertions for protocol compliance", 2)
ASK("Add SystemVerilog Assertions that check protocol rules.",
    "25 SVA properties (52 instances including per-channel copies) in two checker modules, attached with `bind` so the RTL is not modified.",
    "tb/sva/dma_sva.sv")
FIGURE("m3_assert_map.png", "Assertion groups and where they are attached.")
T(["Assertion", "Rule checked"], [
    ["a_apb_access_follows_setup", "a SETUP cycle is always followed by an ACCESS cycle"],
    ["a_apb_stable_in_setup", "PADDR / PWRITE / PWDATA stable from SETUP to ACCESS"],
    ["a_apb_pready, a_apb_slverr_phase", "PREADY in ACCESS; PSLVERR only in ACCESS"],
    ["a_aw_stable, a_ar_stable, a_w_stable", "AXI payload and VALID stable until READY"],
    ["a_aw_attr, a_ar_attr", "size = 4 bytes, INCR burst, word-aligned address"],
    ["a_aw_4k, a_ar_4k", "no burst crosses a 4 KB boundary"],
    ["a_wlast_pos", "WLAST exactly on beat AWLEN+1"],
    ["a_one_phase", "only one of AW / W / AR / B / R active at a time (single outstanding burst)"],
    ["a_gnt_is_req, a_fire_has_gnt", "grant only to a requesting channel; fire only with a valid grant"],
    ["a_prio_order (x4)", "the winner has the highest priority of all requesters"],
    ["a_irq_or", "irq = OR of irq_ch"],
    ["a_sready_only_active", "s_ready only when the channel is an active sensor channel"],
    ["a_no_underflow, a_fifo_bound", "no FIFO pop when empty; level never exceeds 32"],
    ["a_upd_needs_inflight", "completion only for the channel being served"],
    ["a_req_legal, a_beats_range", "request only from an enabled active channel; burst size 1..16"],
    ["a_irq_def", "irq_ch = OR(INT_STAT & INT_EN)"],
    ["a_bad_start_inactive", "an illegal START never activates the channel"],
], [5.6, 11.0], caption="Assertion list")
CODE("""a_aw_4k : assert property (disable iff (!rst_n)
    awvalid |-> ({1'b0, awaddr[11:0]} + ((13'(awlen) + 13'd1) << 2)) <= 13'd4096)
  else $error("AXI: write burst crosses a 4KB boundary");

a_aw_stable : assert property (disable iff (!rst_n)
    awvalid && !awready |=> awvalid && $stable(awaddr) && $stable(awlen))
  else $error("AXI: AW changed/dropped before AWREADY");""", "dma_sva.sv - two examples")
P("Result: **0 assertion failures** in all 13 regression runs. (xsim 2025.2 does not execute `cover property`; the same "
  "scenarios are measured with covergroups instead.)", "j")

# ---- 4.5
H("4.5 Generate functional and code coverage reports", 2)
ASK("Measure coverage and generate functional and code coverage reports.",
    "7 covergroups with crosses, merged over all 13 runs with xcrg; HTML functional and code coverage reports.",
    "tb/uvm/dma_tb_env.svh (dma_coverage), sim/report/functional/, sim/report/code/")
T(["Covergroup", "What it measures"], [
    ["cg_cfg", "channel x mode (memory/sensor) x priority x burst length x transfer length, with crosses"],
    ["cg_axi", "read/write x burst length x response (OKAY/SLVERR/DECERR) x burst ending on a 4 KB boundary"],
    ["cg_arb", "number of simultaneous requesters x winner; winner priority x best competing priority"],
    ["cg_irq", "channel x interrupt cause (DONE / ERR)"],
    ["cg_err", "channel x error code (none, SLVERR, DECERR, bad config)"],
    ["cg_sensor", "channel x word accepted vs back-pressured"],
    ["cg_apb", "register region (global / channel / unmapped) x read/write x PSLVERR"],
], [2.6, 14.0], caption="Covergroups")
FIGURE("m3_cov_chart.png", "Functional coverage per covergroup.")
FIGURE("cov_functional.png", "Merged functional coverage report generated by Vivado xcrg.")
T(["RTL module", "Statement %", "Branch %", "Condition %", "Toggle %"], [
    ["dma_top", "100", "94.1", "83.3", "71.6"], ["dma_channel", "100", "100", "97.2", "54.6"],
    ["dma_sfifo", "100", "100", "100", "99.5"], ["dma_arbiter", "100", "100", "100", "100"],
    ["dma_engine", "100", "100", "97.9", "92.0"],
], [4.6, 3.0, 3.0, 3.0, 3.0], caption="Code coverage of the RTL modules")
FIGURE("cov_code.png", "Code coverage report (Vivado xcrg, per module).")
H("4.5.1 Coverage analysis", 3)
B(["**Functional coverage 100 %.** Every bin of all 7 covergroups was hit at least once over the regression.",
   "**Statement coverage 100 %** on every RTL module: every line of design code was executed.",
   "**Toggle coverage** is lower in dma_top / dma_channel because the high bits of 32-bit addresses and the 24-bit length never "
   "change in a small test memory; this is expected and documented.",
   "Testbench modules (interfaces, assertion checkers, UVM package) appear in the tool totals but are not design code."])
BOX("Milestone 3 - outcome", [
    "All five items are complete: SV testbench (4.1), UVM agents and sequences (4.2), scoreboard and checking (4.3), protocol "
    "assertions (4.4), functional and code coverage reports (4.5).",
    "13/13 runs pass, 0 UVM errors, 0 assertion failures, 100 % functional coverage, 100 % statement coverage."], "E7F6EC", "2F855A")
PAGEBREAK()

# ================================================================ 5. COMPLIANCE
H("5. Milestone compliance checklist", 1)
T(["#", "Milestone item", "Status", "Evidence"], [
    ["M1.1", "Study DMA architecture used in SoCs", "Done", "Section 2.1, docs/01, docs/02"],
    ["M1.2", "Define channel configuration registers", "Done", "Section 2.2, docs/03_register_map.md"],
    ["M1.3", "Design arbitration scheme", "Done", "Section 2.3, Figures 5-6"],
    ["M1.4", "Create datapath and controller architecture", "Done", "Section 2.4, Figures 7-8"],
    ["M1.5", "Develop RTL block diagram", "Done", "Section 2.5, docs/block_diagram.svg"],
    ["M2.1", "Implement DMA channel controller", "Done", "rtl/dma_channel.sv, rtl/dma_sfifo.sv"],
    ["M2.2", "Design arbitration logic", "Done", "rtl/dma_arbiter.sv"],
    ["M2.3", "Implement transfer state machine", "Done", "rtl/dma_engine.sv"],
    ["M2.4", "Integrate interrupt and status registers", "Done", "rtl/dma_channel.sv, rtl/dma_top.sv"],
    ["M2.5", "Perform RTL simulation", "Done", "sim/, sim/logs/ (13/13 pass)"],
    ["M3.1", "Develop SystemVerilog testbench", "Done", "tb/tb_top.sv, tb/if/"],
    ["M3.2", "Create UVM agents and sequences", "Done", "tb/uvm/ (3 agents, 9 sequences)"],
    ["M3.3", "Implement scoreboarding and checking", "Done", "dma_scoreboard, mutation check"],
    ["M3.4", "Add assertions for protocol compliance", "Done", "tb/sva/dma_sva.sv (25 properties)"],
    ["M3.5", "Generate functional and code coverage reports", "Done", "sim/report/ (100 % functional)"],
], [1.4, 6.4, 1.6, 7.2], caption="Every milestone item and its evidence")

H("6. Conclusion", 1)
P("The project delivers a complete, verified multi-channel DMA subsystem for an ADAS platform. Four channels move camera, "
  "radar, LiDAR and diagnostic data into memory with AXI4 bursts while the CPU only programs registers and handles "
  "interrupts. A per-burst arbiter with priority and round-robin keeps the critical radar latency bounded to about one "
  "burst and shares the bus fairly between equal channels. Errors from the memory system and illegal configurations "
  "are detected and isolated to the affected channel.", "j")
P("A UVM environment with an independent reference model, 25 protocol assertions and 7 covergroups verified the design: all "
  "13 regression runs pass, functional coverage is 100 % and every RTL line was executed. The RTL also synthesizes for an "
  "Artix-7 FPGA at 100 MHz.", "j")
H("Possible future work", 2)
B(["anti-starvation aging, so a low-priority channel is guaranteed the bus even under constant high-priority load;",
   "a bus watchdog that reports a memory that never answers;",
   "continuous (ring) mode for back-to-back camera frames without software restarts;",
   "descriptor (scatter-gather) chains in memory for longer, non-contiguous transfers;",
   "multiple outstanding AXI bursts to overlap reads and writes and raise efficiency;",
   "64- or 128-bit data paths for higher camera bandwidth;",
   "a UVM register model (RAL) and formal proofs of the arbiter properties;",
   "place-and-route and on-board testing on an FPGA."])

H("Appendix A. Project files", 1)
T(["Folder / file", "Content"], [
    ["docs/", "requirements, architecture, register map, verification plan, block diagram, this report"],
    ["docs/figures/", "all diagrams, waveforms and coverage screenshots used in the reports"],
    ["rtl/", "dma_pkg, dma_sfifo, dma_channel, dma_arbiter, dma_engine, dma_top"],
    ["tb/tb_top.sv, tb/if/", "testbench top and interfaces"],
    ["tb/sva/", "assertion checkers (bound to the DUT)"],
    ["tb/uvm/", "items, agents, scoreboard, coverage, environment, sequences, tests"],
    ["sim/", "build / run / regression / coverage scripts (.bat and .sh), logs, HTML reports"],
    ["synth/", "Vivado synthesis script, utilisation and timing reports"],
    ["scripts/", "waveform and diagram generators used for the figures"],
], [4.6, 12.0], caption="Folder structure")
H("Appendix B. Glossary", 1)
T(["Term", "Meaning"], [
    ["ADAS", "Advanced Driver Assistance System"], ["DMA", "Direct Memory Access - hardware that copies data without the CPU"],
    ["APB", "Advanced Peripheral Bus - simple control bus"], ["AXI4", "Advanced eXtensible Interface - high-performance burst bus"],
    ["Beat", "one data transfer of a burst (here 4 bytes)"], ["INCR burst", "burst whose address increments each beat"],
    ["W1C", "write-1-to-clear register bit"], ["FIFO", "first-in first-out buffer"],
    ["Round-robin", "taking turns: the next tie goes to the channel after the last winner"], ["UVM", "Universal Verification Methodology"],
    ["SVA", "SystemVerilog Assertions"], ["Scoreboard", "testbench component that compares DUT output with expected results"],
    ["Covergroup", "SystemVerilog construct that counts which scenarios were exercised"],
], [3.2, 13.4], caption="Terms used in this report")

# ---------------------------------------------------------------- header / footer, settings
for s in doc.sections:
    hp = s.header.paragraphs[0]; hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = hp.add_run("ADAS Multi-Channel DMA Controller Subsystem - Milestone Report"); r.font.size = Pt(8.5); r.font.color.rgb = GREY
    fp = s.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run("Page "); r.font.size = Pt(9)
    rr = fp.add_run(); rr.font.size = Pt(9)
    for kind, txt in (("begin", None), ("instr", "PAGE"), ("end", None)):
        if kind == "instr":
            it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = txt; rr._r.append(it)
        else:
            fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), kind); rr._r.append(fc)
s0 = doc.sections[0]
s0.different_first_page_header_footer = True

# pre-filled TOC result (shows headings even before Word updates the field)
last = None
for lvl, txt in TOC_ENTRIES:
    q = TOC_ANCHOR.insert_paragraph_before()
    q.paragraph_format.left_indent = Cm(0 if lvl == 1 else 0.8)
    q.paragraph_format.space_after = Pt(2 if lvl == 2 else 4)
    r = q.add_run(txt); r.font.size = Pt(11 if lvl == 1 else 10)
    if lvl == 1:
        r.bold = True; r.font.color.rgb = NAVY
    last = q
fe = OxmlElement("w:fldChar"); fe.set(qn("w:fldCharType"), "end")
er = last.add_run(); er._r.append(fe)

upd = OxmlElement("w:updateFields"); upd.set(qn("w:val"), "true")
se = doc.settings.element
nxt = None
for tag in ("hdrShapeDefaults", "footnotePr", "endnotePr", "compat", "docVars", "rsids", "mathPr", "attachedSchema",
            "themeFontLang", "clrSchemeMapping", "doNotIncludeSubdocsInStats", "doNotAutoCompressPictures", "forceUpgrade",
            "captions", "readModeInkLockDown", "smartTagType", "schemaLibrary", "shapeDefaults", "doNotEmbedSmartTags",
            "decimalSymbol", "listSeparator"):
    nxt = se.find(qn("w:" + tag))
    if nxt is not None:
        break
nxt.addprevious(upd) if nxt is not None else se.append(upd)
z = se.find(qn("w:zoom"))
if z is not None and z.get(qn("w:percent")) is None:
    z.set(qn("w:percent"), "100")
doc.core_properties.title = "ADAS Multi-Channel DMA Controller Subsystem - Milestone Report"
doc.core_properties.subject = "L&T EduTech - VLSI Principles, RTL Design and Verification - Project 1"

doc.save(OUT)
print("wrote", OUT, "figures:", fig_no[0], "tables:", tab_no[0])
